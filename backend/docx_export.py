#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docx_export.py — turn a report's HTML (report_builder / wgs_report) into an
editable Word document.

    from docx_export import html_to_docx
    html_to_docx(html_string, "report.docx")

The PDF and the Word file come from the same HTML, so they always say the same
thing; this module only maps the small set of tags those builders write onto Word:

    h1 -> Title,  h2 -> Heading 1,  h3 -> Heading 2   (so Word's navigation pane works)
    p / ul / li / table / div.note / div.kpis / p.cmd
    figure: the inline SVG is rasterised to PNG (cairosvg, or rsvg-convert as a
            fallback) and placed at the width it has in the PDF; figcaption and the
            "Folder:" line follow as paragraphs.

Everything stays editable: headings, tables and captions are real Word objects,
not pictures of them. Only the charts are images.

Requires python-docx and cairosvg (pip install python-docx cairosvg) in the
environment the backend runs in.
"""

from __future__ import annotations

import html as _html
import io
import re
import shutil
import subprocess
from html.parser import HTMLParser

__version__ = "1.0.0"

INK, ACC, MUT, FAINT = "16242B", "1E6E8C", "4E6672", "7D949E"
RULE, NOTE_BG, NOTE_ACC_BG, HEAD_BG = "D6E0E4", "FAF6EC", "EDF4F7", "EDF4F7"
WARN = "96660A"
VOID = {"br", "img", "meta", "hr", "input", "link", "col", "wbr"}
BLOCK = {"h1", "h2", "h3", "h4", "p", "div", "ul", "ol", "li", "table", "figure",
         "figcaption", "caption", "section", "header", "footer", "pre", "blockquote"}


class MissingDependency(RuntimeError):
    pass


# ── a tiny DOM ────────────────────────────────────────────────────────────────
class Node:
    __slots__ = ("tag", "attrs", "children", "parent")

    def __init__(self, tag, attrs=None, parent=None):
        self.tag, self.attrs, self.children, self.parent = tag, dict(attrs or {}), [], parent

    @property
    def classes(self):
        return set((self.attrs.get("class") or "").split())

    def text(self):
        out = []
        for c in self.children:
            out.append(c if isinstance(c, str) else c.text())
        return "".join(out)

    def find(self, tag):
        for c in self.children:
            if isinstance(c, Node):
                if c.tag == tag:
                    return c
                f = c.find(tag)
                if f is not None:
                    return f
        return None

    def find_all(self, tag):
        out = []
        for c in self.children:
            if isinstance(c, Node):
                if c.tag == tag:
                    out.append(c)
                out += c.find_all(tag)
        return out


class _Builder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.cur = self.root
        self.style = []          # <style> contents are dropped
        self.title = ""
        self._in = None

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script", "title"):
            self._in = tag
            return
        n = Node(tag, attrs, self.cur)
        self.cur.children.append(n)
        if tag not in VOID:
            self.cur = n

    def handle_startendtag(self, tag, attrs):
        n = Node(tag, attrs, self.cur)
        self.cur.children.append(n)

    def handle_endtag(self, tag):
        if tag in ("style", "script", "title"):
            self._in = None
            return
        if tag in VOID:
            return
        n = self.cur
        while n is not None and n.tag != tag:
            n = n.parent
        if n is not None and n.parent is not None:
            self.cur = n.parent

    def handle_data(self, data):
        if self._in == "title":
            self.title += data
            return
        if self._in:
            return
        self.cur.children.append(data)


# ── SVG → PNG ─────────────────────────────────────────────────────────────────
_XML_ENT = {"amp", "lt", "gt", "quot", "apos"}


def _svg_clean(svg: str):
    """Fixed pixel size from the viewBox, an xmlns, and no HTML-only entities."""
    m = re.search(r'viewBox="\s*[-\d.]+\s+[-\d.]+\s+([\d.]+)\s+([\d.]+)"', svg)
    w, h = (float(m.group(1)), float(m.group(2))) if m else (520.0, 300.0)
    head_end = svg.index(">")
    head = svg[:head_end]
    head = re.sub(r'\swidth="[^"]*"', "", head)
    head = re.sub(r'\sheight="[^"]*"', "", head)
    head = re.sub(r'\sstyle="[^"]*"', "", head)
    if "xmlns=" not in head:
        head += ' xmlns="http://www.w3.org/2000/svg"'
    head += f' width="{w:g}" height="{h:g}"'
    svg = head + svg[head_end:]
    svg = re.sub(r"<(?=[^A-Za-z/!?])", "&lt;", svg)     # a bare "<" in text (e.g. "<.1")
    svg = re.sub(r"&([A-Za-z][A-Za-z0-9]*);",
                 lambda mm: mm.group(0) if mm.group(1) in _XML_ENT
                 else "&#%d;" % ord(_html.unescape(mm.group(0))[0]), svg)
    return svg, w, h


def svg_to_png(svg: str, scale: float = 3.0):
    """(png bytes, width px, height px) at `scale` × the viewBox size."""
    svg, w, h = _svg_clean(svg)
    try:
        import cairosvg  # noqa: WPS433
        png = cairosvg.svg2png(bytestring=svg.encode("utf-8"),
                               output_width=int(w * scale), output_height=int(h * scale),
                               background_color="white")
        return png, w, h
    except ImportError:
        pass
    exe = shutil.which("rsvg-convert")
    if exe:
        p = subprocess.run([exe, "-w", str(int(w * scale)), "-h", str(int(h * scale)),
                            "-b", "white", "-f", "png"],
                           input=svg.encode("utf-8"), capture_output=True, timeout=60)
        if p.returncode == 0 and p.stdout:
            return p.stdout, w, h
    raise MissingDependency("cairosvg")


# ── Word helpers ──────────────────────────────────────────────────────────────
def _docx():
    try:
        import docx  # noqa: F401
        return docx
    except ImportError as e:
        raise MissingDependency("python-docx") from e


def _rgb(hexs):
    from docx.shared import RGBColor
    return RGBColor.from_string(hexs)


def _shade(cell_or_par, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _cell_borders(cell, **edges):
    """edges: top=(size_eighths, color) etc.; None removes the edge."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge in ("top", "left", "bottom", "right"):
        v = edges.get(edge, "keep")
        if v == "keep":
            continue
        e = OxmlElement(f"w:{edge}")
        if v is None:
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(v[0]))
            e.set(qn("w:color"), v[1])
        b.append(e)


def _table_no_borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def _repeat_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def _field(paragraph, instr):
    """A Word field (PAGE, NUMPAGES) inside a paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"),
                       ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            run._r.append(f)
        elif text == instr:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = f" {instr} "
            run._r.append(t)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
    return run


def _ws(s: str) -> str:
    return re.sub(r"\s+", " ", s)


# ── the converter ─────────────────────────────────────────────────────────────
class _Writer:
    def __init__(self, doc, svgs, content_w_cm):
        self.doc, self.svgs, self.cw = doc, svgs, content_w_cm
        self.pending_break = False
        self.after_block = False

    # paragraphs with inline formatting
    def para(self, container=None, style=None):
        tgt = container if container is not None else self.doc
        p = tgt.add_paragraph(style=style) if style else tgt.add_paragraph()
        if self.pending_break and container is None:
            p.paragraph_format.page_break_before = True
            self.pending_break = False
        if self.after_block and container is None:
            # space after a table / note box, without an empty paragraph that can
            # spill onto a page of its own
            p.paragraph_format.space_before = _pt(10)
            self.after_block = False
        return p

    def inline(self, p, node, fmt=None):
        fmt = dict(fmt or {})
        if isinstance(node, str):
            t = _ws(node)
            if not t:
                return
            # no leading space at the start of a paragraph
            if not p.runs and not p.text:
                t = t.lstrip()
                if not t:
                    return
            r = p.add_run(t)
            self._apply(r, fmt)
            return
        tag, cls = node.tag, node.classes
        if tag == "br":
            p.add_run().add_break()
            return
        if tag == "svgref":
            self.picture(p, node)
            return
        if tag in ("b", "strong"):
            fmt["bold"] = True
        elif tag in ("i", "em"):
            fmt["italic"] = True
        elif tag in ("code", "tt", "kbd"):
            fmt["mono"] = True
        elif tag == "sup":
            fmt["sup"] = True
        elif tag == "sub":
            fmt["sub"] = True
        if "mono" in cls:
            fmt["mono"] = True
        if "dim" in cls:
            fmt["color"] = FAINT
        if "hl" in cls:
            fmt["bold"], fmt["color"] = True, ACC
        st = node.attrs.get("style") or ""
        m = re.search(r"font-size:\s*([\d.]+)pt", st)
        if m:
            fmt["size"] = float(m.group(1))
        m = re.search(r"(?:^|;)\s*color:\s*#([0-9A-Fa-f]{6})", st)
        if m:
            fmt["color"] = m.group(1).upper()
        for c in node.children:
            if isinstance(c, Node) and c.tag in BLOCK:
                # a block inside inline content (e.g. <div> in a table cell): keep it inline
                self.inline(p, c, fmt)
            else:
                self.inline(p, c, fmt)

    @staticmethod
    def _apply(r, fmt):
        from docx.shared import Pt
        if fmt.get("bold"):
            r.bold = True
        if fmt.get("italic"):
            r.italic = True
        if fmt.get("mono"):
            r.font.name = "Consolas"
        if fmt.get("sup"):
            r.font.superscript = True
        if fmt.get("sub"):
            r.font.subscript = True
        if fmt.get("color"):
            r.font.color.rgb = _rgb(fmt["color"])
        if fmt.get("size"):
            r.font.size = Pt(fmt["size"])

    def picture(self, p, node):
        from docx.shared import Cm
        i = int(node.attrs.get("data-i", "-1"))
        if not (0 <= i < len(self.svgs)):
            return
        png, w, _h = svg_to_png(self.svgs[i])
        width = Cm(min(self.cw, self.cw * w / 520.0))
        p.add_run().add_picture(io.BytesIO(png), width=width)

    # blocks
    def block(self, node, container=None):
        if isinstance(node, str):
            if _ws(node).strip():
                p = self.para(container)
                self.inline(p, node)
            return
        tag, cls = node.tag, node.classes
        if "pbreak" in cls and container is None:
            self.pending_break = True
        if tag in ("h1", "h2", "h3", "h4"):
            self.heading(node, container)
        elif tag == "p":
            self.paragraph(node, container)
        elif tag in ("ul", "ol"):
            style = "List Bullet" if tag == "ul" else "List Number"
            for li in node.children:
                if isinstance(li, Node) and li.tag == "li":
                    p = self.para(container, style)
                    for c in li.children:
                        self.inline(p, c)
                    if "refs" in cls:
                        for r in p.runs:
                            r.font.size = _pt(8.5)
        elif tag == "table":
            self.table(node, container)
        elif tag == "figure":
            self.figure(node, container)
        elif tag == "div" and "kpis" in cls:
            self.kpis(node, container)
        elif tag == "div" and "note" in cls:
            self.note(node, container, accent="acc" in cls)
        elif tag == "div" and "eyebrow" in cls:
            p = self.para(container)
            r = p.add_run(_ws(node.text()).strip().upper())
            r.font.size, r.font.color.rgb = _pt(7.5), _rgb(ACC)
            r.font.name = "Consolas"
        elif tag == "svgref":
            p = self.para(container)
            self.picture(p, node)
        elif tag in ("div", "section", "header", "footer", "body", "html", "root",
                     "blockquote", "main", "article"):
            self.mixed(node, container)
        elif tag == "figcaption" or tag == "caption":
            p = self.para(container)
            self.inline(p, node, {"size": 8.5, "color": MUT})
        else:
            p = self.para(container)
            self.inline(p, node)

    def mixed(self, node, container):
        """Children that mix inline content and blocks: inline runs share a paragraph."""
        p = None
        for c in node.children:
            if isinstance(c, Node) and (c.tag in BLOCK or c.tag == "svgref"):
                p = None
                self.block(c, container)
            else:
                if isinstance(c, str) and not _ws(c).strip():
                    continue
                if p is None:
                    p = self.para(container)
                self.inline(p, c)

    def heading(self, node, container):
        lvl = {"h1": 0, "h2": 1, "h3": 2, "h4": 3}[node.tag]
        text = _ws(node.text()).strip()
        if container is not None:
            p = container.add_paragraph()
            r = p.add_run(text)
            r.bold = True
            return
        h = self.doc.add_heading(text, level=lvl)
        self.after_block = False
        if self.pending_break:
            h.paragraph_format.page_break_before = True
            self.pending_break = False
        h.paragraph_format.keep_with_next = True

    def paragraph(self, node, container):
        cls = node.classes
        p = self.para(container)
        if "cmd" in cls:
            _shade(p, "F1F5F6")
            self.inline(p, node, {"mono": True, "size": 8})
            return
        if "tcap" in cls:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = _pt(8)
            self.inline(p, node, {"size": 9})
            return
        if "sub" in cls:
            self.inline(p, node, {"color": MUT, "size": 11})
            return
        if "project" in cls:
            self.inline(p, node, {"color": ACC, "bold": True, "size": 14})
            return
        self.inline(p, node)

    def kpis(self, node, container):
        boxes = [c for c in node.children if isinstance(c, Node) and "kpi" in c.classes]
        if not boxes:
            return
        tgt = container if container is not None else self.doc
        t = tgt.add_table(rows=1, cols=len(boxes))
        _table_no_borders(t)
        for cell, b in zip(t.rows[0].cells, boxes):
            v = next((c for c in b.children if isinstance(c, Node) and "v" in c.classes), None)
            l_ = next((c for c in b.children if isinstance(c, Node) and "l" in c.classes), None)
            _cell_borders(cell, top=(16, ACC), left=(4, RULE), right=(4, RULE), bottom=(4, RULE))
            p = cell.paragraphs[0]
            r = p.add_run(_ws(v.text()).strip() if v else "")
            r.bold, r.font.size, r.font.name = True, _pt(14), "Consolas"
            p2 = cell.add_paragraph()
            r2 = p2.add_run(_ws(l_.text()).strip() if l_ else "")
            r2.font.size, r2.font.color.rgb = _pt(8), _rgb(MUT)
        self.after_block = container is None

    def note(self, node, container, accent=False):
        tgt = container if container is not None else self.doc
        t = tgt.add_table(rows=1, cols=1)
        _table_no_borders(t)
        cell = t.rows[0].cells[0]
        _shade(cell, NOTE_ACC_BG if accent else NOTE_BG)
        _cell_borders(cell, left=(24, ACC if accent else WARN))
        first = cell.paragraphs[0]
        used = False
        for c in node.children:
            if isinstance(c, Node) and c.tag == "p":
                p = first if not used else cell.add_paragraph()
                used = True
                self.inline(p, c, {"size": 9.5})
            elif isinstance(c, Node):
                self.block(c, cell)
            elif _ws(c).strip():
                p = first if not used else cell.add_paragraph()
                used = True
                self.inline(p, c, {"size": 9.5})
        if container is None:
            self.after_block = True

    def table(self, node, container):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        head = node.find("thead")
        body = node.find("tbody")
        hrows = head.find_all("tr") if head else []
        brows = body.find_all("tr") if body else [r for r in node.find_all("tr") if r not in hrows]
        cells_of = lambda tr: [c for c in tr.children if isinstance(c, Node) and c.tag in ("th", "td")]
        rows = hrows + brows
        if not rows:
            return
        ncol = max(len(cells_of(r)) for r in rows)
        tgt = container if container is not None else self.doc
        if self.pending_break and container is None:
            pb = self.doc.add_paragraph()
            pb.paragraph_format.page_break_before = True
            self.pending_break = False
        t = tgt.add_table(rows=len(rows), cols=ncol)
        t.style = self.doc.styles["Table Grid"]
        _table_no_borders(t)
        # Column widths from content (Word's autofit squeezes long text columns into
        # a sliver), and right alignment only for columns that hold numbers.
        texts = [[_ws(c.text()).strip() for c in cells_of(tr)] for tr in rows]
        num_re = re.compile(r"^[\d,.\s%<>=+\-\u2013\u2014/×]*$")
        need, numeric = [], []
        for ci in range(ncol):
            body = [r[ci] for r, tr in zip(texts, rows) if tr not in hrows and ci < len(r)]
            headw = max((max((len(w) for w in r[ci].split()), default=0)
                         for r, tr in zip(texts, rows) if tr in hrows and ci < len(r)), default=0)
            longest = max((len(x) for x in body), default=0)
            numeric.append(bool(body) and sum(1 for x in body if num_re.match(x)) >= 0.8 * len(body))
            chars = max(longest, min(headw, 7 if numeric[-1] else 10), 3)
            need.append(0.19 * min(chars, 30) + 0.45)          # cm at 9 pt, plus cell padding
        from docx.shared import Cm
        avail = 17.4 if container is None else 15.5
        if sum(need) > avail:
            # take the excess from the wide text columns first, never below 2.2 cm
            excess = sum(need) - avail
            wide = [i for i in range(ncol) if not numeric[i] and need[i] > 2.2]
            room = sum(need[i] - 2.2 for i in wide)
            if room > 0:
                cut = min(excess, room)
                for i in wide:
                    need[i] -= cut * (need[i] - 2.2) / room
            if sum(need) > avail:
                k = avail / sum(need)
                need = [x * k for x in need]
        else:
            k = avail / sum(need)
            need = [x * k for x in need]
        weights, tot = need, avail
        t.autofit = False
        for ci in range(ncol):
            wcm = Cm(avail * weights[ci] / tot)
            t.columns[ci].width = wcm            # tblGrid: what LibreOffice reads
            for row in t.rows:
                row.cells[ci].width = wcm        # tcW: what Word reads
        for ri, tr in enumerate(rows):
            is_head = tr in hrows
            if is_head:
                _repeat_header(t.rows[ri])
            for ci, src in enumerate(cells_of(tr)):
                cell = t.cell(ri, ci)
                p = cell.paragraphs[0]
                st = (src.attrs.get("style") or "") + " ".join(
                    (c.attrs.get("style") or "") for c in src.children if isinstance(c, Node))
                left = ci == 0 or not numeric[ci] or "text-align:left" in st.replace(" ", "")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.RIGHT
                if is_head:
                    self.inline(p, src, {"bold": True, "size": 8, "color": MUT})
                    _shade(cell, HEAD_BG)
                    _cell_borders(cell, bottom=(8, INK))
                else:
                    self.inline(p, src, {"size": 9})
                    _cell_borders(cell, bottom=(4, RULE))
        # small tables move to the next page whole instead of leaving a header behind
        from docx.oxml import OxmlElement
        for ri, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if len(t.rows) <= 15 and ri < len(t.rows) - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        cap = node.find("caption")
        if cap is not None:
            p = self.para(container)
            self.inline(p, cap, {"size": 8.5, "color": FAINT})
        elif container is None:
            self.after_block = True

    def figure(self, node, container):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        for c in node.children:
            if isinstance(c, Node) and c.tag == "svgref":
                p = self.para(container)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                self.picture(p, c)
            elif isinstance(c, Node) and c.tag == "figcaption":
                p = self.para(container)
                self.inline(p, c, {"size": 9})
            elif isinstance(c, Node) and "folder" in c.classes:
                p = self.para(container)
                self.inline(p, c, {"bold": True, "size": 9})
            elif isinstance(c, Node):
                self.block(c, container)


def _pt(v):
    from docx.shared import Pt
    return Pt(v)


def _setup(doc, title, footer_label):
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(2.0)
    st = doc.styles
    normal = st["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (("Title", 22, INK), ("Heading 1", 15, INK),
                              ("Heading 2", 12, INK), ("Heading 3", 11, INK)):
        s = st[name]
        s.font.name, s.font.size, s.font.bold = "Cambria", Pt(size), True
        s.font.color.rgb = _rgb(color)
    st["Heading 1"].paragraph_format.space_before = Pt(14)
    st["Heading 2"].paragraph_format.space_before = Pt(10)
    # header: report title, right; footer: job label left, page x / y right
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(title)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(8), _rgb(MUT)
    # footer as a borderless two-cell table: tab stops are rendered differently
    # by Word and LibreOffice, a table is not
    ft = sec.footer.add_table(rows=1, cols=2, width=Cm(17.4))
    _table_no_borders(ft)
    ft.autofit = False
    ft.columns[0].width = ft.rows[0].cells[0].width = Cm(13.4)
    ft.columns[1].width = ft.rows[0].cells[1].width = Cm(4.0)
    lp = ft.rows[0].cells[0].paragraphs[0]
    r = lp.add_run(footer_label)
    r.font.size, r.font.color.rgb, r.font.name = Pt(7.5), _rgb(FAINT), "Consolas"
    rp = ft.rows[0].cells[1].paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for part in ("PAGE", " / ", "NUMPAGES"):
        rr = rp.add_run(part) if part.strip() == "/" else _field(rp, part)
        rr.font.size, rr.font.color.rgb = Pt(7.5), _rgb(FAINT)
    empty = sec.footer.paragraphs[0]
    if not empty.text:
        empty._p.getparent().remove(empty._p)


def html_to_docx(html_doc: str, out_path, title: str | None = None,
                 footer_label: str = "") -> str:
    """Write `html_doc` as a .docx at `out_path`. Raises MissingDependency when
    python-docx or an SVG rasteriser is not installed."""
    docx = _docx()
    svgs = []

    def grab(m):
        svgs.append(m.group(0))
        return f'<svgref data-i="{len(svgs) - 1}"></svgref>'

    body = re.sub(r"<svg\b.*?</svg>", grab, html_doc, flags=re.S | re.I)
    b = _Builder()
    b.feed(body)
    b.close()
    if svgs:                       # fail early and clearly, before a half-built file
        svg_to_png('<svg viewBox="0 0 10 10"><rect width="1" height="1"/></svg>', 1)
    doc = docx.Document()
    t = title or _ws(b.title).strip() or "Report"
    if not footer_label:
        m = re.search(r'@bottom-left\s*\{\s*content:\s*"([^"]*)"', html_doc)
        footer_label = m.group(1) if m else ""
    _setup(doc, t, footer_label)
    w = _Writer(doc, svgs, content_w_cm=16.5)
    body_node = b.root.find("body") or b.root
    w.mixed(body_node, None)
    # python-docx starts with one empty paragraph; drop it if still empty
    first = doc.paragraphs[0] if doc.paragraphs else None
    if first is not None and not first.text and not first.runs:
        first._p.getparent().remove(first._p)
    doc.core_properties.title = t
    doc.save(str(out_path))
    return str(out_path)


if __name__ == "__main__":
    import sys
    src, dst = sys.argv[1], sys.argv[2]
    html_to_docx(open(src, encoding="utf-8").read(), dst)
    print(dst)
